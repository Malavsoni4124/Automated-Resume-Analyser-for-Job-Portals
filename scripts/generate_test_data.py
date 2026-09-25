import os
import json
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
import docx

RESUME_DIR = "tests/test_data/resumes"
TRUTH_DIR = "tests/test_data/ground_truth"

os.makedirs(RESUME_DIR, exist_ok=True)
os.makedirs(TRUTH_DIR, exist_ok=True)

def save_truth(filename: str, data: dict):
    with open(os.path.join(TRUTH_DIR, f"{filename}.json"), "w") as f:
        json.dump(data, f, indent=2)

def generate_standard_pdf():
    name = "standard_resume"
    c = canvas.Canvas(os.path.join(RESUME_DIR, f"{name}.pdf"), pagesize=letter)
    
    textobject = c.beginText(50, 750)
    textobject.textLine("John Doe")
    textobject.textLine("john.doe@example.com | (555) 123-4567 | github.com/johndoe")
    textobject.textLine("")
    textobject.textLine("SUMMARY")
    textobject.textLine("Experienced Software Engineer.")
    textobject.textLine("")
    textobject.textLine("EXPERIENCE")
    textobject.textLine("Google | Software Engineer | Jan 2020 - Present")
    textobject.textLine("- Developed scalable backend services.")
    textobject.textLine("")
    textobject.textLine("EDUCATION")
    textobject.textLine("MIT | B.S. Computer Science | 2016 - 2020")
    textobject.textLine("")
    textobject.textLine("SKILLS")
    textobject.textLine("Python, Java, Machine Learning, AWS")
    
    c.drawText(textobject)
    c.save()
    
    truth = {
        "contact": {"name": "John Doe", "email": "john.doe@example.com", "phone": "(555) 123-4567", "links": ["github.com/johndoe"]},
        "summary": "Experienced Software Engineer.",
        "experience": [{"company": "Google", "role": "Software Engineer"}],
        "education": [{"institution": "MIT", "degree": "B.S. Computer Science"}],
        "skills": ["Python", "Java", "Machine Learning", "AWS"]
    }
    save_truth(name, truth)

def generate_scanned_pdf():
    # Simulate scanned PDF by putting almost no text layer, or drawing it as an image.
    # To keep it simple, we just put very little text.
    name = "scanned_image"
    c = canvas.Canvas(os.path.join(RESUME_DIR, f"{name}.pdf"), pagesize=letter)
    c.drawString(100, 100, " ") # Near empty text
    c.save()
    
    truth = {"status": "requires_ocr"}
    save_truth(name, truth)

def generate_skills_only_docx():
    name = "skills_only"
    doc = docx.Document()
    doc.add_heading('Jane Smith', 0)
    doc.add_paragraph('jane.smith@email.com')
    doc.add_heading('Technical Skills', level=1)
    doc.add_paragraph('React, JavaScript, Docker, Kubernetes')
    doc.save(os.path.join(RESUME_DIR, f"{name}.docx"))
    
    truth = {
        "contact": {"name": "Jane Smith", "email": "jane.smith@email.com"},
        "skills": ["React", "JavaScript", "Docker", "Kubernetes"]
    }
    save_truth(name, truth)

def generate_dense_no_headers_pdf():
    name = "dense_no_headers"
    c = canvas.Canvas(os.path.join(RESUME_DIR, f"{name}.pdf"), pagesize=letter)
    textobject = c.beginText(50, 750)
    textobject.textLine("Alice Walker alice@walker.net")
    textobject.textLine("I am a data scientist who worked at Meta from 2021 to 2023 doing SQL and Python.")
    textobject.textLine("I studied at Stanford University.")
    c.drawText(textobject)
    c.save()
    
    truth = {
        "contact": {"name": "Alice Walker", "email": "alice@walker.net"},
        "experience": [{"company": "Meta"}],
        "education": [{"institution": "Stanford University"}],
        "skills": ["SQL", "Python"]
    }
    save_truth(name, truth)

if __name__ == "__main__":
    generate_standard_pdf()
    generate_scanned_pdf()
    generate_skills_only_docx()
    generate_dense_no_headers_pdf()
    
def generate_real_validation_set():
    # Real 1: highly unstructured
    name = "real_1_unstructured"
    doc = docx.Document()
    doc.add_paragraph("Michael Chang\nSoftware Architect\nmichael.c@example.org | 555-987-6543")
    doc.add_paragraph("Work:\nSenior Engineer at Amazon (2018-2022). Promoted to Architect in 2022. I worked heavily with Python and Kubernetes.")
    doc.add_paragraph("Education: MS Computer Science from UC Berkeley.")
    doc.save(os.path.join(RESUME_DIR, f"{name}.docx"))
    save_truth(name, {
        "category": "validation",
        "contact": {"name": "Michael Chang", "email": "michael.c@example.org", "phone": "555-987-6543"},
        "experience": [{"company": "Amazon", "role": "Senior Engineer"}, {"company": "Amazon", "role": "Architect"}],
        "education": [{"institution": "UC Berkeley", "degree": "MS Computer Science"}],
        "skills": ["Python", "Kubernetes"]
    })

    # Real 2: messy formatting PDF
    name = "real_2_messy_pdf"
    c = canvas.Canvas(os.path.join(RESUME_DIR, f"{name}.pdf"), pagesize=letter)
    textobject = c.beginText(50, 750)
    textobject.textLine("S A R A   J O N E S")
    textobject.textLine("Data Scientist @ Netflix | sara.jones@netflix.com")
    textobject.textLine("Ed: PhD Physics, MIT")
    textobject.textLine("Proficient in Machine Learning, SQL, Python.")
    c.drawText(textobject)
    c.save()
    save_truth(name, {
        "category": "validation",
        "contact": {"name": "Sara Jones", "email": "sara.jones@netflix.com"},
        "experience": [{"company": "Netflix", "role": "Data Scientist"}],
        "education": [{"institution": "MIT", "degree": "PhD Physics"}],
        "skills": ["Machine Learning", "SQL", "Python"]
    })
    
    # Real 3, 4, 5... (generate simple placeholders for real set)
    for i in range(3, 6):
        name = f"real_{i}_template"
        doc = docx.Document()
        doc.add_paragraph(f"Candidate {i}\ncandidate{i}@email.com")
        doc.add_paragraph(f"Developer at Startup {i}")
        doc.add_paragraph(f"BSc from University {i}")
        doc.save(os.path.join(RESUME_DIR, f"{name}.docx"))
        save_truth(name, {
            "category": "validation",
            "contact": {"name": f"Candidate {i}", "email": f"candidate{i}@email.com"},
            "experience": [{"company": f"Startup {i}", "role": "Developer"}],
            "education": [{"institution": f"University {i}", "degree": "BSc"}],
            "skills": []
        })

if __name__ == "__main__":
    generate_standard_pdf()
    generate_scanned_pdf()
    generate_skills_only_docx()
    generate_dense_no_headers_pdf()
    generate_real_validation_set()
    
    for i in range(1, 11):
        name = f"variant_standard_{i}"
        doc = docx.Document()
        doc.add_paragraph(f"Test User {i}")
        doc.add_paragraph(f"user{i}@test.com")
        doc.add_heading("Skills", level=1)
        doc.add_paragraph("Python, AWS")
        doc.save(os.path.join(RESUME_DIR, f"{name}.docx"))
        
        truth = {
            "category": "standard",
            "contact": {"name": f"Test User {i}", "email": f"user{i}@test.com"},
            "skills": ["Python", "AWS"]
        }
        save_truth(name, truth)
        
    print("Test data generated successfully.")
